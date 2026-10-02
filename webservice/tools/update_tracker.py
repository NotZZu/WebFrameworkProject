"""시험 관리 엑셀(테스트관리.xlsx) 생성/갱신.

- 케이스 정의(@case 레지스트리, cases.json) → '케이스관리' 시트(우선순위·화면·비고는 사람이 직접 관리, 갱신 시 보존)
- pytest JUnit XML 파싱 → '케이스별결과'(최신) 시트 덮어쓰기 + '실행이력' 시트에 한 줄 추가
- '대시보드' 시트: 수식 + 차트로 개발/테스트 진척도 시각화 (엑셀에서 열면 자동 계산)

사용:
  python tools/update_tracker.py --run                       # pytest 실행 → XML → 엑셀 갱신
  python tools/update_tracker.py --xml report.xml --cases cases.json   # (CI) 이미 있는 결과 사용
옵션: --out tracker/테스트관리.xlsx
참고: openpyxl 은 기존 차트를 보존하지 못하므로 매번 통합문서를 새로 만들고, 수동 입력값과 이력만 이전 파일에서 이어받는다.
"""
import argparse
import datetime
import json
import os
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "tracker" / "테스트관리.xlsx"
FONT = "Arial"
NAVY, ACCENT = "1E2761", "3D5AFE"
C_PASS, C_RED, C_FAIL, C_SKIP = "2E7D32", "EF9A9A", "B71C1C", "9E9E9E"
SCREENS = [("SC-01", "로그인·회원가입"), ("SC-02", "대시보드"), ("SC-03", "가명·익명 처리"),
           ("SC-04", "CAVP 신청"), ("SC-05", "신청 조회"), ("SC-06", "검증서 이력"),
           ("SC-07", "위변조 탐지"), ("공통", "공통·초기데이터·E2E")]
REQ_SCREEN = {"FR-01": "SC-01", "FR-02": "SC-03", "FR-03": "SC-04", "FR-04": "SC-06", "FR-05": "SC-07", "FR-14": "공통"}
REQ_PRIORITY = {"FR-01": "Must", "FR-02": "Must", "FR-03": "Must", "FR-04": "Should", "FR-05": "Should", "FR-14": "Should"}
KIND_KO = {"positive": "정상", "negative": "비정상", "boundary": "경계", "security": "보안"}
LEVEL_KO = {"UT": "단위", "IT": "통합", "ST": "시스템"}
CASE_HDR_ROW, CASE_FIRST = 4, 5
RES_HDR_ROW, RES_FIRST = 3, 4
HIS_HDR_ROW, HIS_FIRST = 3, 4
MAXR = 600
KST = datetime.timezone(datetime.timedelta(hours=9))


# ---------------------------------------------------------------- 입력 파싱
def natural_key(case_id: str):
    order = {"ST": 0, "IT": 1, "UT": 2}
    return (order.get(case_id[:2], 9), [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", case_id)])


def default_screen(c) -> str:
    cid = c["id"]
    if cid.startswith("ST-01") or cid.startswith("UT-SEED") or cid in ("IT-11", "IT-12", "IT-31", "IT-32"):
        return "SC-01" if cid.startswith("ST-01") else "공통"
    return REQ_SCREEN.get(c["req"], "공통")


def parse_junit(xml_path):
    """→ ({(module, func): {pass, fail, red, skip, reason}}, 총 pytest 건수, 소요 초)"""
    root = ET.parse(xml_path).getroot()
    suite = root if root.tag == "testsuite" else root.find("testsuite")
    per = {}
    runs = 0
    for tc in suite.iter("testcase"):
        runs += 1
        key = (tc.get("classname"), tc.get("name").split("[")[0])
        r = per.setdefault(key, dict(**{"pass": 0, "fail": 0, "red": 0, "skip": 0}, reason=""))
        bad = tc.find("failure") if tc.find("failure") is not None else tc.find("error")
        sk = tc.find("skipped")
        if bad is not None:
            r["fail"] += 1
            r["reason"] = r["reason"] or (bad.get("message") or "").splitlines()[0][:150]
        elif sk is not None and "TDD Red" in (sk.get("message") or ""):
            r["red"] += 1
        elif sk is not None:
            r["skip"] += 1
        else:
            r["pass"] += 1
    return per, runs, float(suite.get("time") or 0)


def collect_results(cases, per):
    res = {}
    for c in cases:
        r = per.get((c["module"], c["func"].split(".")[-1]))
        if r:
            res[c["id"]] = r
    return res


def git(*a):
    try:
        return subprocess.run(["git", *a], capture_output=True, text=True, cwd=ROOT).stdout.strip()
    except OSError:
        return ""


# ---------------------------------------------------------------- 기존 파일에서 이어받기
def load_previous(path):
    manual, history = {}, []
    if Path(path).exists():
        wb = load_workbook(path)
        if "케이스관리" in wb.sheetnames:
            ws = wb["케이스관리"]
            for row in ws.iter_rows(min_row=CASE_FIRST, values_only=True):
                if row[0]:
                    manual[row[0]] = {"screen": row[3], "priority": row[7], "note": row[8]}
        if "실행이력" in wb.sheetnames:
            ws = wb["실행이력"]
            for row in ws.iter_rows(min_row=HIS_FIRST, max_col=10, values_only=True):
                if row[0]:
                    history.append(list(row))
    return manual, history


# ---------------------------------------------------------------- 스타일 헬퍼
def f(**kw):
    return Font(name=FONT, **{"size": 10, **kw})


THIN = Side(style="thin", color="DFE3EE")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = f(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = f(bold=True, size=16, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(size=9, color="6B7488", italic=True)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def cell(ws, ref, value, **style):
    c = ws[ref]
    c.value = value
    c.font = style.pop("font", f())
    c.border = BOX
    for k, v in style.items():
        setattr(c, k, v)
    return c


def status_formats(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'],
        fill=PatternFill("solid", bgColor="C8E6C9", fgColor="C8E6C9"), font=Font(name=FONT, color="1B5E20", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"RED"'],
        fill=PatternFill("solid", bgColor="FFE0E0", fgColor="FFE0E0"), font=Font(name=FONT, color="C62828", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'],
        fill=PatternFill("solid", bgColor="B71C1C", fgColor="B71C1C"), font=Font(name=FONT, color="FFFFFF", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"SKIP"'],
        fill=PatternFill("solid", bgColor="EEEEEE", fgColor="EEEEEE"), font=Font(name=FONT, color="616161")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"미실행"'],
        fill=PatternFill("solid", bgColor="FFF8E1", fgColor="FFF8E1"), font=Font(name=FONT, color="8D6E00")))


# ---------------------------------------------------------------- 통합문서 생성
def build_workbook(cases, results, history, manual, out_path):
    cases = sorted(cases, key=lambda c: natural_key(c["id"]))
    wb = Workbook()
    dash = wb.active
    dash.title = "대시보드"
    cs = wb.create_sheet("케이스관리")
    rs = wb.create_sheet("케이스별결과")
    hs = wb.create_sheet("실행이력")

    # ---- 케이스관리
    title(cs, "시험 케이스 관리", "노란색 칸(화면·우선순위·비고)만 직접 수정하세요. 나머지는 테스트 코드(@case)에서 자동 동기화되며, 수정 내용은 갱신해도 유지됩니다. 예: 화면=SC-01, 우선순위=Must")
    cs["A3"] = "※ '현재 상태'는 '케이스별결과' 시트에서 자동 계산 — PASS 통과 · RED 구현 전(기대된 실패) · FAIL 실패(병합 차단) · SKIP 건너뜀 · 미실행"
    cs["A3"].font = f(size=9, color="6B7488")
    header(cs, CASE_HDR_ROW, ["케이스 ID", "구분", "요구사항", "화면", "시험 유형", "시험 항목", "기대결과", "우선순위", "비고", "현재 상태"])
    input_fill = PatternFill("solid", fgColor="FFFDE7")
    for i, c in enumerate(cases):
        r = CASE_FIRST + i
        m = manual.get(c["id"], {})
        vals = [c["id"], LEVEL_KO.get(c["id"][:2], "기타"), c["req"], m.get("screen") or default_screen(c),
                KIND_KO.get(c["kind"], c["kind"]), c["title"], c["expect"],
                m.get("priority") or REQ_PRIORITY.get(c["req"], "Could"), m.get("note")]
        for j, v in enumerate(vals, 1):
            x = cs.cell(row=r, column=j, value=v)
            x.font = f(color="0000FF") if j in (4, 8, 9) else f()
            x.border = BOX
            x.alignment = Alignment(vertical="top", wrap_text=j in (6, 7, 9))
            if j in (4, 8, 9):
                x.fill = input_fill
        s = cs.cell(row=r, column=10,
                    value=f"=IFERROR(INDEX('케이스별결과'!$B${RES_FIRST}:$B${MAXR},MATCH($A{r},'케이스별결과'!$A${RES_FIRST}:$A${MAXR},0)),\"미실행\")")
        s.font = f(bold=True); s.border = BOX; s.alignment = Alignment(horizontal="center", vertical="top")
    last_case = CASE_FIRST + len(cases) - 1
    widths(cs, [13, 7, 9, 9, 9, 46, 52, 10, 28, 11])
    cs.freeze_panes = cs.cell(row=CASE_FIRST, column=2)
    cs.auto_filter.ref = f"A{CASE_HDR_ROW}:J{last_case}"
    status_formats(cs, f"J{CASE_FIRST}:J{max(last_case, CASE_FIRST)}")
    dv1 = DataValidation(type="list", formula1="=대시보드!$A$24:$A$31", allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Must,Should,Could"', allow_blank=True)
    cs.add_data_validation(dv1); cs.add_data_validation(dv2)
    dv1.add(f"D{CASE_FIRST}:D{MAXR}"); dv2.add(f"H{CASE_FIRST}:H{MAXR}")

    # ---- 케이스별결과 (최신 실행 결과, 자동 생성)
    title(rs, "케이스별 최신 결과 (자동 생성 — 수정하지 마세요)", "pytest JUnit XML 을 파싱한 값입니다. 결과(B)는 건수로부터 수식으로 판정합니다: 실패>0 → FAIL, Red>0 → RED, 통과>0 → PASS")
    header(rs, RES_HDR_ROW, ["케이스 ID", "결과", "실행 건수", "통과", "실패", "Red", "건너뜀", "실패 사유(첫 줄)"])
    for i, c in enumerate(cases):
        r = RES_FIRST + i
        x = results.get(c["id"])
        if x is None:
            continue
        rs.cell(row=r, column=1, value=c["id"])
        rs.cell(row=r, column=2, value=f'=IF(E{r}>0,"FAIL",IF(F{r}>0,"RED",IF(D{r}>0,"PASS","SKIP")))')
        rs.cell(row=r, column=3, value=f"=SUM(D{r}:G{r})")
        for col, k in ((4, "pass"), (5, "fail"), (6, "red"), (7, "skip")):
            rs.cell(row=r, column=col, value=x[k])
        rs.cell(row=r, column=8, value=x["reason"] or None)
        for col in range(1, 9):
            rs.cell(row=r, column=col).font = f(bold=(col == 2))
            rs.cell(row=r, column=col).border = BOX
            if col in (2, 3, 4, 5, 6, 7):
                rs.cell(row=r, column=col).alignment = Alignment(horizontal="center")
    widths(rs, [13, 9, 10, 7, 7, 7, 8, 80])
    rs.freeze_panes = "B4"
    status_formats(rs, f"B{RES_FIRST}:B{MAXR}")

    # ---- 실행이력
    title(hs, "실행 이력 (실행할 때마다 한 줄 추가)", "케이스수·PASS·RED·FAIL·SKIP 은 해당 실행의 케이스(ID) 단위 집계(고정 기록값), Green% = PASS ÷ (PASS+RED+FAIL)")
    header(hs, HIS_HDR_ROW, ["실행 일시", "브랜치", "커밋", "케이스 수", "PASS", "RED", "FAIL", "SKIP", "pytest 건수", "소요(초)", "Green %"])
    for i, h in enumerate(history[-MAXR + 10:]):
        r = HIS_FIRST + i
        for j, v in enumerate(h[:10], 1):
            x = hs.cell(row=r, column=j, value=v)
            x.font = f(color="0000FF" if j >= 4 else "000000"); x.border = BOX
            x.alignment = Alignment(horizontal="center" if j != 2 else "left")
        g = hs.cell(row=r, column=11, value=f"=IFERROR(E{r}/(E{r}+F{r}+G{r}),0)")
        g.number_format = "0.0%"; g.font = f(bold=True); g.border = BOX; g.alignment = Alignment(horizontal="center")
    last_h = HIS_FIRST + max(len(history), 1) - 1
    widths(hs, [18, 26, 10, 10, 8, 8, 8, 8, 12, 10, 10])
    hs.freeze_panes = "A4"
    hs.conditional_formatting.add(f"K{HIS_FIRST}:K{MAXR}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="66BB6A"))

    # ---- 대시보드
    title(dash, "시험·개발 진척도 대시보드", "자동화 검증시스템 · TDD(Red→Green) 진척 — 모든 수치는 다른 시트에서 수식으로 계산됩니다 (엑셀에서 열면 자동 갱신)")
    dash["A3"] = f'="최종 갱신: "&IFERROR(INDEX(\'실행이력\'!$A${HIS_FIRST}:$A${MAXR},COUNTA(\'실행이력\'!$A${HIS_FIRST}:$A${MAXR})),"-")&"  ·  브랜치: "&IFERROR(INDEX(\'실행이력\'!$B${HIS_FIRST}:$B${MAXR},COUNTA(\'실행이력\'!$B${HIS_FIRST}:$B${MAXR})),"-")'
    dash["A3"].font = f(size=10, bold=True, color=ACCENT)
    rng = f"'케이스관리'!$J${CASE_FIRST}:$J${MAXR}"
    kp = [("전체 케이스", f"=COUNTA('케이스관리'!$A${CASE_FIRST}:$A${MAXR})", NAVY, "0"),
          ("통과 (Green)", f'=COUNTIF({rng},"PASS")', C_PASS, "0"),
          ("구현 전 (Red)", f'=COUNTIF({rng},"RED")', "C62828", "0"),
          ("실패 (병합 차단)", f'=COUNTIF({rng},"FAIL")', C_FAIL, "0"),
          ("Green 진행률", "=IFERROR(C6/(C6+E6+G6),0)", ACCENT, "0.0%")]
    for k, (lab, formula, color, fmt) in enumerate(kp):
        c1 = 1 + 2 * k
        a, b = get_column_letter(c1), get_column_letter(c1 + 1)
        dash.merge_cells(f"{a}5:{b}5"); dash.merge_cells(f"{a}6:{b}6")
        dash[f"{a}5"] = lab
        dash[f"{a}5"].font = f(size=9, bold=True, color="FFFFFF")
        dash[f"{a}5"].fill = PatternFill("solid", fgColor=color)
        dash[f"{a}5"].alignment = Alignment(horizontal="center")
        dash[f"{a}6"] = formula
        dash[f"{a}6"].font = f(size=22, bold=True, color=color)
        dash[f"{a}6"].alignment = Alignment(horizontal="center", vertical="center")
        dash[f"{a}6"].number_format = fmt
        dash[f"{b}5"].fill = PatternFill("solid", fgColor=color)
    dash.row_dimensions[6].height = 38
    # KPI 수식이 병합 셀 첫 칸(A,C,E,G,I)을 가리키도록: 통과 C6, Red E6, 실패 G6
    # 상태 분포
    dash["A8"] = "상태 분포"; dash["A8"].font = f(bold=True, size=12, color=NAVY)
    header(dash, 9, ["상태", "설명", "건수", "비율"])
    st_rows = [("PASS", "통과", C_PASS), ("RED", "구현 전(기대된 실패)", C_RED), ("FAIL", "실패", C_FAIL), ("SKIP", "건너뜀", C_SKIP), ("미실행", "결과 없음", "FFE082")]
    for i, (code, desc, _) in enumerate(st_rows):
        r = 10 + i
        cell(dash, f"A{r}", code, font=f(bold=True)); cell(dash, f"B{r}", desc)
        cell(dash, f"C{r}", f'=COUNTIF({rng},A{r})', alignment=Alignment(horizontal="center"))
        cell(dash, f"D{r}", f"=IFERROR(C{r}/SUM($C$10:$C$14),0)", alignment=Alignment(horizontal="center"), number_format="0.0%")
    status_formats(dash, "A10:A14")
    pie = DoughnutChart(); pie.title = "시험 케이스 상태"; pie.holeSize = 55
    pie.add_data(Reference(dash, min_col=3, min_row=9, max_row=14), titles_from_data=True)
    pie.set_categories(Reference(dash, min_col=1, min_row=10, max_row=14))
    for idx, (_, _, col) in enumerate(st_rows):
        pt = DataPoint(idx=idx); pt.graphicalProperties.solidFill = col; pie.series[0].dPt.append(pt)
    pie.height, pie.width = 7.2, 11
    dash.add_chart(pie, "F8")

    # 화면별 진척
    dash["A22"] = "화면별 진척"; dash["A22"].font = f(bold=True, size=12, color=NAVY)
    header(dash, 23, ["화면", "화면명", "전체", "통과", "Red", "실패", "진척(통과%)"])
    scr_rng = f"'케이스관리'!$D${CASE_FIRST}:$D${MAXR}"
    for i, (code, name) in enumerate(SCREENS):
        r = 24 + i
        cell(dash, f"A{r}", code, font=f(bold=True)); cell(dash, f"B{r}", name)
        cell(dash, f"C{r}", f"=COUNTIF({scr_rng},A{r})", alignment=Alignment(horizontal="center"))
        for col, code_ in (("D", "PASS"), ("E", "RED"), ("F", "FAIL")):
            cell(dash, f"{col}{r}", f'=COUNTIFS({scr_rng},$A{r},{rng},"{code_}")', alignment=Alignment(horizontal="center"))
        cell(dash, f"G{r}", f'=IF(C{r}=0,"케이스 없음",D{r}/C{r})', alignment=Alignment(horizontal="center"), number_format="0%")
    dash.conditional_formatting.add("G24:G31", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="66BB6A"))
    bar = BarChart(); bar.type = "col"; bar.grouping = "stacked"; bar.overlap = 100
    bar.title = "화면별 통과 / Red / 실패"
    for col, color in ((4, C_PASS), (5, C_RED), (6, C_FAIL)):
        bar.add_data(Reference(dash, min_col=col, min_row=23, max_row=31), titles_from_data=True)
        bar.series[-1].graphicalProperties.solidFill = color
    bar.set_categories(Reference(dash, min_col=1, min_row=24, max_row=31))
    bar.height, bar.width = 6.2, 15
    dash.add_chart(bar, "I22")

    # 요구사항별 / 구분별
    dash["A34"] = "요구사항별 진척"; dash["A34"].font = f(bold=True, size=12, color=NAVY)
    header(dash, 35, ["요구사항", "우선순위", "전체", "통과", "Red", "실패", "진척(통과%)"])
    req_rng = f"'케이스관리'!$C${CASE_FIRST}:$C${MAXR}"
    reqs = ["FR-01", "FR-02", "FR-03", "FR-04", "FR-05", "FR-14"]
    for i, fr in enumerate(reqs):
        r = 36 + i
        cell(dash, f"A{r}", fr, font=f(bold=True)); cell(dash, f"B{r}", REQ_PRIORITY[fr], alignment=Alignment(horizontal="center"))
        cell(dash, f"C{r}", f"=COUNTIF({req_rng},A{r})", alignment=Alignment(horizontal="center"))
        for col, code_ in (("D", "PASS"), ("E", "RED"), ("F", "FAIL")):
            cell(dash, f"{col}{r}", f'=COUNTIFS({req_rng},$A{r},{rng},"{code_}")', alignment=Alignment(horizontal="center"))
        cell(dash, f"G{r}", f'=IF(C{r}=0,"케이스 없음",D{r}/C{r})', alignment=Alignment(horizontal="center"), number_format="0%")
    dash.conditional_formatting.add("G36:G41", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="66BB6A"))
    dash["I34"] = "시험 단계별 진척"; dash["I34"].font = f(bold=True, size=12, color=NAVY)
    header(dash, 35, ["단계", "전체", "통과", "Red", "실패"], col=9)
    lvl_rng = f"'케이스관리'!$B${CASE_FIRST}:$B${MAXR}"
    for i, lv in enumerate(["시스템", "통합", "단위"]):
        r = 36 + i
        cell(dash, f"I{r}", lv, font=f(bold=True))
        cell(dash, f"J{r}", f"=COUNTIF({lvl_rng},I{r})", alignment=Alignment(horizontal="center"))
        for col, code_ in (("K", "PASS"), ("L", "RED"), ("M", "FAIL")):
            cell(dash, f"{col}{r}", f'=COUNTIFS({lvl_rng},$I{r},{rng},"{code_}")', alignment=Alignment(horizontal="center"))

    # 추이
    dash["A44"] = "진척 추이 (실행할 때마다 갱신)"; dash["A44"].font = f(bold=True, size=12, color=NAVY)
    line = LineChart(); line.title = "실행별 PASS / RED / FAIL 케이스 수 (Red → Green 전환)"
    for col, color in ((5, C_PASS), (6, "E57373"), (7, C_FAIL)):
        line.add_data(Reference(hs, min_col=col, min_row=HIS_HDR_ROW, max_row=max(last_h, HIS_FIRST + 1)), titles_from_data=True)
        s = line.series[-1]; s.graphicalProperties.line.solidFill = color; s.graphicalProperties.line.width = 28000
        s.marker.symbol = "circle"; s.marker.size = 7
        s.marker.graphicalProperties.solidFill = color; s.marker.graphicalProperties.line.solidFill = color
    line.set_categories(Reference(hs, min_col=1, min_row=HIS_FIRST, max_row=max(last_h, HIS_FIRST + 1)))
    line.height, line.width = 8, 26
    line.y_axis.title = "케이스 수"; line.y_axis.delete = False; line.x_axis.delete = False
    dash.add_chart(line, "A45")

    widths(dash, [12, 22, 9, 9, 9, 9, 14, 3, 12, 9, 9, 9, 9, 9])
    dash.sheet_view.showGridLines = False
    dash.sheet_properties.tabColor = ACCENT
    cs.sheet_properties.tabColor = "FFC107"
    for w in (dash, cs, rs, hs):          # 인쇄/PDF: 가로, 1페이지 너비에 맞춤
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.calculation.fullCalcOnLoad = True
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return wb


def make_history_row(results, cases, runs, secs, branch, commit, when=None):
    cnt = {"PASS": 0, "RED": 0, "FAIL": 0, "SKIP": 0}
    for c in cases:
        x = results.get(c["id"])
        if not x:
            continue
        cnt["FAIL" if x["fail"] else "RED" if x["red"] else "PASS" if x["pass"] else "SKIP"] += 1
    when = when or datetime.datetime.now(KST).strftime("%Y-%m-%d %H:%M")
    return [when, branch, commit, sum(cnt.values()), cnt["PASS"], cnt["RED"], cnt["FAIL"], cnt["SKIP"], runs, round(secs, 1)]


def update(xml_path, cases_path, out_path, branch=None, commit=None):
    cases = json.load(open(cases_path, encoding="utf-8"))
    per, runs, secs = parse_junit(xml_path)
    results = collect_results(cases, per)
    manual, history = load_previous(out_path)
    branch = branch or os.environ.get("GITHUB_REF_NAME") or git("rev-parse", "--abbrev-ref", "HEAD") or "-"
    commit = commit or (os.environ.get("GITHUB_SHA") or git("rev-parse", "HEAD") or "-")[:7]
    row = make_history_row(results, cases, runs, secs, branch, commit)
    if history and history[-1][1] == branch and history[-1][2] == commit:
        history[-1] = row           # 같은 커밋 재실행은 덮어쓰기
    else:
        history.append(row)
    build_workbook(cases, results, history, manual, out_path)
    return row


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xml"); ap.add_argument("--cases"); ap.add_argument("--out", default=str(DEFAULT_OUT))
    ap.add_argument("--run", action="store_true", help="pytest 를 실행해 XML 을 새로 만든다")
    a = ap.parse_args()
    xml = a.xml or str(ROOT / "report.xml")
    cases = a.cases or str(ROOT / "cases.json")
    if a.run or not a.xml:
        subprocess.run([sys.executable, "-m", "pytest", "-q", f"--junitxml={xml}"], cwd=ROOT)
    if a.run or not a.cases:
        subprocess.run([sys.executable, str(ROOT / "scripts" / "export_cases.py"), cases], cwd=ROOT, check=True)
    row = update(xml, cases, a.out)
    print(f"갱신 완료 → {a.out}\n  {row}")


if __name__ == "__main__":
    main()
