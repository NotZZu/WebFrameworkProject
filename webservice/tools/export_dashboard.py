"""시험 관리 엑셀 → GitHub Pages 용 정적 대시보드(HTML) 변환.

사용: python tools/export_dashboard.py [--xlsx tracker/테스트관리.xlsx] [--out ../assets/tracker]
출력: <out>/dashboard.html (외부 라이브러리 없는 단일 파일, iframe 임베드용), <out>/test-tracker.xlsx (다운로드용 사본)
엑셀의 '입력값'(케이스관리 · 케이스별결과 건수 · 실행이력)만 읽어 상태를 직접 계산하므로 엑셀 수식 재계산이 필요 없다.
"""
import argparse
import html
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCREEN_NAMES = {"SC-01": "로그인·회원가입", "SC-02": "대시보드", "SC-03": "가명·익명 처리", "SC-04": "CAVP 신청",
                "SC-05": "신청 조회", "SC-06": "검증서 이력", "SC-07": "위변조 탐지", "공통": "공통·초기데이터·E2E"}
REQ_ORDER = ["FR-01", "FR-02", "FR-03", "FR-04", "FR-05", "FR-14"]


def status_of(r):
    if r is None:
        return "미실행"
    p, f, red, sk = r
    return "FAIL" if f else "RED" if red else "PASS" if p else "SKIP"


def read(xlsx):
    wb = load_workbook(xlsx)
    res = {}
    ws = wb["케이스별결과"]
    for row in ws.iter_rows(min_row=4, max_col=8, values_only=True):
        if row[0]:
            res[row[0]] = ((row[3] or 0), (row[4] or 0), (row[5] or 0), (row[6] or 0), row[7])
    cases = []
    for row in wb["케이스관리"].iter_rows(min_row=5, max_col=9, values_only=True):
        if not row[0]:
            continue
        r = res.get(row[0])
        cases.append(dict(id=row[0], level=row[1], req=row[2], screen=row[3], kind=row[4], title=row[5],
                          expect=row[6], priority=row[7], note=row[8] or "",
                          status=status_of(r[:4] if r else None), reason=(r[4] if r else "") or ""))
    hist = []
    for row in wb["실행이력"].iter_rows(min_row=4, max_col=10, values_only=True):
        if row[0]:
            hist.append(dict(when=str(row[0]), branch=row[1], commit=row[2], total=row[3], p=row[4], red=row[5], f=row[6], sk=row[7]))
    return cases, hist


def count(cases, key, val):
    sub = [c for c in cases if c[key] == val]
    return dict(total=len(sub), p=sum(c["status"] == "PASS" for c in sub), red=sum(c["status"] == "RED" for c in sub),
                f=sum(c["status"] == "FAIL" for c in sub))


def donut(p, red, f, sk):
    tot = max(p + red + f + sk, 1)
    segs, off = [], 0.0
    for v, col in ((p, "#2e7d32"), (red, "#ef9a9a"), (f, "#b71c1c"), (sk, "#9e9e9e")):
        if v:
            ln = 100 * v / tot
            segs.append(f'<circle r="15.9155" cx="21" cy="21" fill="none" stroke="{col}" stroke-width="6" '
                        f'stroke-dasharray="{ln:.2f} {100 - ln:.2f}" stroke-dashoffset="{-off:.2f}" transform="rotate(-90 21 21)"/>')
            off += ln
    pct = round(100 * p / max(p + red + f, 1))
    return (f'<svg viewBox="0 0 42 42" width="190" height="190"><circle r="15.9155" cx="21" cy="21" fill="none" stroke="#eef0f6" stroke-width="6"/>'
            + "".join(segs) + f'<text x="21" y="22.5" text-anchor="middle" font-size="7" font-weight="700" fill="#1e2761">{pct}%</text>'
            f'<text x="21" y="28" text-anchor="middle" font-size="2.6" fill="#6b7488">Green</text></svg>')


def bars(rows):
    out = []
    for label, d in rows:
        t = max(d["total"], 1)
        seg = "".join(f'<i style="width:{100 * d[k] / t:.2f}%;background:{c}" title="{n} {d[k]}"></i>'
                      for k, c, n in (("p", "#2e7d32", "통과"), ("red", "#ef9a9a", "Red"), ("f", "#b71c1c", "실패")))
        pct = f'{round(100 * d["p"] / d["total"])}%' if d["total"] else "케이스 없음"
        out.append(f'<div class="brow"><span class="bl">{html.escape(label)}</span><div class="bar">{seg}</div>'
                   f'<span class="bn">{d["p"]}/{d["total"]} · {pct}</span></div>')
    return "".join(out)


def trend(hist):
    if len(hist) < 2:
        return '<p class="muted">실행 이력이 2건 이상 쌓이면 추이 그래프가 표시됩니다.</p>'
    W, H, L, B, T = 760, 250, 46, 38, 16
    ymax = max(max(h["total"] or 0, 1) for h in hist)
    n = len(hist)
    x = lambda i: L + (W - L - 20) * i / (n - 1)
    y = lambda v: T + (H - T - B) * (1 - v / ymax)
    g = "".join(f'<line x1="{L}" x2="{W - 20}" y1="{y(v):.1f}" y2="{y(v):.1f}" stroke="#e6e9f2"/>'
                f'<text x="{L - 6}" y="{y(v) + 3:.1f}" text-anchor="end" font-size="10" fill="#6b7488">{v}</text>'
                for v in sorted({0, round(ymax / 2), ymax}))
    lines = ""
    for key, col in (("p", "#2e7d32"), ("red", "#ef9a9a"), ("f", "#b71c1c")):
        pts = " ".join(f"{x(i):.1f},{y(h[key] or 0):.1f}" for i, h in enumerate(hist))
        dots = "".join(f'<circle cx="{x(i):.1f}" cy="{y(h[key] or 0):.1f}" r="4" fill="{col}"><title>{html.escape(h["when"])} · {h[key]}</title></circle>'
                       for i, h in enumerate(hist))
        lines += f'<polyline points="{pts}" fill="none" stroke="{col}" stroke-width="2.5"/>{dots}'
    step = max(1, n // 8)
    xl = "".join(f'<text x="{x(i):.1f}" y="{H - 16}" text-anchor="middle" font-size="9" fill="#6b7488">{html.escape(h["when"][5:16] if h["when"][:2] == "20" and len(h["when"]) >= 16 else h["when"][:10])}</text>'
                 for i, h in enumerate(hist) if i % step == 0 or i == n - 1)
    return f'<svg viewBox="0 0 {W} {H}" width="100%">{g}{lines}{xl}</svg>'


def render(cases, hist):
    cnt = {k: sum(c["status"] == k for c in cases) for k in ("PASS", "RED", "FAIL", "SKIP", "미실행")}
    green = round(100 * cnt["PASS"] / max(cnt["PASS"] + cnt["RED"] + cnt["FAIL"], 1), 1)
    last = hist[-1] if hist else {"when": "-", "branch": "-", "commit": "-"}
    scr = [(f"{k} {SCREEN_NAMES.get(k, '')}", count(cases, "screen", k)) for k in SCREEN_NAMES]
    req = [(k, count(cases, "req", k)) for k in REQ_ORDER]
    lvl = [(k, count(cases, "level", k)) for k in ("시스템", "통합", "단위")]
    data = json.dumps(cases, ensure_ascii=False).replace("</", "<\\/")
    return f"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>시험·개발 진척도</title><style>
:root{{--navy:#1e2761;--accent:#3d5afe;--muted:#6b7488;--line:#dfe3ee;--bg:#f4f6fa}}
*{{box-sizing:border-box}}body{{margin:0;font-family:system-ui,-apple-system,"Noto Sans KR","Malgun Gothic",sans-serif;color:var(--navy);background:var(--bg);padding:20px}}
h1{{margin:0;font-size:22px}}.sub{{color:var(--muted);font-size:13px;margin:4px 0 16px}}
.top{{display:flex;flex-wrap:wrap;justify-content:space-between;gap:10px;align-items:flex-end}}
.dl{{background:var(--accent);color:#fff;text-decoration:none;font-size:13px;font-weight:700;padding:8px 14px;border-radius:999px}}
.kpis{{display:grid;grid-template-columns:repeat(auto-fit,minmax(140px,1fr));gap:12px;margin:12px 0 18px}}
.kpi{{background:#fff;border-radius:12px;padding:14px 16px;border-top:4px solid var(--c);box-shadow:0 2px 8px rgba(30,39,97,.06)}}
.kpi b{{display:block;font-size:30px;color:var(--c)}}.kpi span{{font-size:12px;color:var(--muted)}}
.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(340px,1fr));gap:16px}}
.card{{background:#fff;border-radius:12px;padding:16px 18px;box-shadow:0 2px 8px rgba(30,39,97,.06)}}
.card h2{{margin:0 0 12px;font-size:15px}}.donutwrap{{display:flex;gap:18px;align-items:center;flex-wrap:wrap}}
.legend div{{font-size:13px;margin:5px 0}}.legend i{{display:inline-block;width:11px;height:11px;border-radius:3px;margin-right:7px;vertical-align:-1px}}
.brow{{display:grid;grid-template-columns:130px 1fr 110px;gap:8px;align-items:center;margin:7px 0;font-size:12.5px}}
.bar{{display:flex;height:14px;background:#eef0f6;border-radius:7px;overflow:hidden}}.bar i{{display:block}}.bn{{color:var(--muted);text-align:right;font-size:12px}}
.muted{{color:var(--muted);font-size:13px}}.full{{grid-column:1/-1}}
.filters{{display:flex;flex-wrap:wrap;gap:8px;margin-bottom:10px}}.filters select,.filters input{{padding:7px 9px;border:1px solid var(--line);border-radius:8px;font:inherit;font-size:13px}}
table{{border-collapse:collapse;width:100%;font-size:12.5px}}th{{background:var(--navy);color:#fff;padding:7px 8px;text-align:left;position:sticky;top:0}}
td{{border-bottom:1px solid var(--line);padding:6px 8px;vertical-align:top}}.tw{{max-height:520px;overflow:auto;border:1px solid var(--line);border-radius:8px}}
.st{{font-weight:700;border-radius:6px;padding:2px 8px;display:inline-block;font-size:11.5px}}
.PASS{{background:#c8e6c9;color:#1b5e20}}.RED{{background:#ffe0e0;color:#c62828}}.FAIL{{background:#b71c1c;color:#fff}}.SKIP{{background:#eee;color:#616161}}.미실행{{background:#fff8e1;color:#8d6e00}}
</style></head><body>
<div class="top"><div><h1>시험·개발 진척도</h1><div class="sub">TDD(Red → Green) 진척 · 최종 갱신 {html.escape(last['when'])} · 브랜치 {html.escape(str(last['branch']))} · 커밋 {html.escape(str(last['commit']))}</div></div>
<a class="dl" href="test-tracker.xlsx" download>엑셀(테스트관리.xlsx) 다운로드</a></div>
<div class="kpis">
<div class="kpi" style="--c:#1e2761"><b>{len(cases)}</b><span>전체 케이스</span></div>
<div class="kpi" style="--c:#2e7d32"><b>{cnt['PASS']}</b><span>통과 (Green)</span></div>
<div class="kpi" style="--c:#c62828"><b>{cnt['RED']}</b><span>구현 전 (Red)</span></div>
<div class="kpi" style="--c:#b71c1c"><b>{cnt['FAIL']}</b><span>실패 (병합 차단)</span></div>
<div class="kpi" style="--c:#3d5afe"><b>{green}%</b><span>Green 진행률</span></div></div>
<div class="grid">
<div class="card"><h2>상태 분포</h2><div class="donutwrap">{donut(cnt['PASS'], cnt['RED'], cnt['FAIL'], cnt['SKIP'])}<div class="legend">
<div><i style="background:#2e7d32"></i>통과 {cnt['PASS']}</div><div><i style="background:#ef9a9a"></i>구현 전(Red) {cnt['RED']}</div><div><i style="background:#b71c1c"></i>실패 {cnt['FAIL']}</div><div><i style="background:#9e9e9e"></i>건너뜀 {cnt['SKIP']}</div></div></div></div>
<div class="card"><h2>화면별 진척</h2>{bars(scr)}</div>
<div class="card"><h2>요구사항별 진척</h2>{bars(req)}<h2 style="margin-top:18px">시험 단계별 진척</h2>{bars(lvl)}</div>
<div class="card full"><h2>실행별 PASS / RED / FAIL 추이</h2>{trend(hist)}
<div class="legend" style="display:flex;gap:16px"><div><i style="background:#2e7d32"></i>PASS</div><div><i style="background:#ef9a9a"></i>RED</div><div><i style="background:#b71c1c"></i>FAIL</div></div></div>
<div class="card full"><h2>케이스 목록 ({len(cases)}건)</h2><div class="filters">
<select id="fs"><option value="">화면 전체</option></select><select id="fk"><option value="">상태 전체</option><option>PASS</option><option>RED</option><option>FAIL</option><option>SKIP</option><option>미실행</option></select>
<select id="fl"><option value="">단계 전체</option><option>시스템</option><option>통합</option><option>단위</option></select><input id="fq" placeholder="검색 (ID·항목)" size="22"></div>
<div class="tw"><table><thead><tr><th>ID</th><th>단계</th><th>요구</th><th>화면</th><th>유형</th><th>시험 항목</th><th>우선</th><th>상태</th></tr></thead><tbody id="tb"></tbody></table></div></div>
</div>
<script>
const D={data};const $=id=>document.getElementById(id);
[...new Set(D.map(c=>c.screen))].sort().forEach(s=>$("fs").add(new Option(s,s)));
const esc=s=>String(s??"").replace(/[&<>"]/g,m=>({{"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}}[m]));
function draw(){{const s=$("fs").value,k=$("fk").value,l=$("fl").value,q=$("fq").value.toLowerCase();
$("tb").innerHTML=D.filter(c=>(!s||c.screen==s)&&(!k||c.status==k)&&(!l||c.level==l)&&(!q||(c.id+c.title).toLowerCase().includes(q))).map(c=>
`<tr><td><b>${{esc(c.id)}}</b></td><td>${{esc(c.level)}}</td><td>${{esc(c.req)}}</td><td>${{esc(c.screen)}}</td><td>${{esc(c.kind)}}</td><td title="${{esc(c.expect)}}">${{esc(c.title)}}${{c.reason?'<br><small style="color:#b71c1c">'+esc(c.reason)+'</small>':''}}</td><td>${{esc(c.priority)}}</td><td><span class="st ${{c.status}}">${{c.status}}</span></td></tr>`).join("")}}
["fs","fk","fl"].forEach(i=>$(i).onchange=draw);$("fq").oninput=draw;draw();
</script></body></html>"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=str(ROOT / "tracker" / "테스트관리.xlsx"))
    ap.add_argument("--out", default=str(ROOT.parent / "assets" / "tracker"))
    a = ap.parse_args()
    cases, hist = read(a.xlsx)
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    (out / "dashboard.html").write_text(render(cases, hist), encoding="utf-8")
    shutil.copyfile(a.xlsx, out / "test-tracker.xlsx")
    print(f"→ {out / 'dashboard.html'} ({len(cases)} cases, {len(hist)} runs)")


if __name__ == "__main__":
    main()
