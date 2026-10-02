"""도구 시험: tools/update_tracker.py (시험 관리 엑셀 생성·갱신). 시험계획서 케이스(@case)가 아닌 개발 도구 자체의 검증."""
import importlib.util
import json
import pathlib

import pytest
from openpyxl import load_workbook

_spec = importlib.util.spec_from_file_location(
    "update_tracker", pathlib.Path(__file__).resolve().parents[2] / "tools" / "update_tracker.py")
ut = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(ut)

CASES = [
    {"id": "UT-AUTH-01", "req": "FR-01", "title": "해시", "expect": "bcrypt", "kind": "positive", "func": "test_a", "module": "tests.unit.m"},
    {"id": "UT-AUTH-02", "req": "FR-01", "title": "검증", "expect": "True", "kind": "negative", "func": "test_b", "module": "tests.unit.m"},
    {"id": "UT-PRIV-01", "req": "FR-02", "title": "가명", "expect": "PSN", "kind": "boundary", "func": "test_c", "module": "tests.unit.m"},
]
XML = """<?xml version="1.0"?><testsuites><testsuite tests="4" failures="1" skipped="1" time="1.5">
<testcase classname="tests.unit.m" name="test_a" time="0.1"/>
<testcase classname="tests.unit.m" name="test_b[x]" time="0.1"><failure message="assert 1 == 2">trace</failure></testcase>
<testcase classname="tests.unit.m" name="test_b[y]" time="0.1"/>
<testcase classname="tests.unit.m" name="test_c" time="0.1"><skipped type="pytest.xfail" message="TDD Red: UT-PRIV-01 구현 전"/></testcase>
</testsuite></testsuites>"""


@pytest.fixture()
def files(tmp_path):
    (tmp_path / "r.xml").write_text(XML, encoding="utf-8")
    (tmp_path / "c.json").write_text(json.dumps(CASES), encoding="utf-8")
    return tmp_path / "r.xml", tmp_path / "c.json", tmp_path / "t.xlsx"


def test_parse_classifies_pass_fail_red(files):
    per, runs, _ = ut.parse_junit(files[0])
    res = ut.collect_results(CASES, per)
    assert runs == 4
    assert res["UT-AUTH-01"]["pass"] == 1
    assert res["UT-AUTH-02"]["fail"] == 1 and res["UT-AUTH-02"]["pass"] == 1 and "assert 1 == 2" in res["UT-AUTH-02"]["reason"]
    assert res["UT-PRIV-01"]["red"] == 1


def test_update_creates_sheets_and_history_row(files):
    xml, cases, out = files
    row = ut.update(str(xml), str(cases), str(out), branch="feature/x", commit="abc1234")
    assert row[3:8] == [3, 1, 1, 1, 0]          # 케이스 3: PASS 1 · RED 1 · FAIL 1 (FAIL 판정이 우선)
    wb = load_workbook(out)
    assert wb.sheetnames == ["대시보드", "케이스관리", "케이스별결과", "실행이력"]
    assert wb["케이스관리"]["A5"].value == "UT-AUTH-01"
    assert str(wb["케이스별결과"]["B4"].value).startswith("=IF(")      # 판정은 수식
    assert str(wb["대시보드"]["C6"].value).startswith("=COUNTIF(")      # KPI 는 수식


def test_update_preserves_manual_columns_and_appends_history(files):
    xml, cases, out = files
    ut.update(str(xml), str(cases), str(out), branch="b1", commit="c1")
    wb = load_workbook(out)
    ws = wb["케이스관리"]
    ws["D5"], ws["H5"], ws["I5"] = "SC-02", "Could", "직접 적은 메모"
    wb.save(out)
    ut.update(str(xml), str(cases), str(out), branch="b1", commit="c2")      # 다른 커밋 → 이력 추가
    ut.update(str(xml), str(cases), str(out), branch="b1", commit="c2")      # 같은 커밋 → 덮어쓰기
    wb = load_workbook(out)
    ws = wb["케이스관리"]
    assert (ws["D5"].value, ws["H5"].value, ws["I5"].value) == ("SC-02", "Could", "직접 적은 메모")
    hist = [r for r in wb["실행이력"].iter_rows(min_row=4, max_col=3, values_only=True) if r[0]]
    assert [h[2] for h in hist] == ["c1", "c2"]
